import re
from io import BytesIO

from docx import Document
from docx.document import Document as DocumentoDocx
from docx.text.paragraph import Paragraph

_NEGRITO = re.compile(r"\*\*(.+?)\*\*")
_NUMERADO = re.compile(r"^(\d+)\\?\.\s+(.*)$")
_MARCADOR = re.compile(r"^-\s+(.*)$")


def _adicionar_runs(paragrafo: Paragraph, texto: str) -> None:
    posicao = 0
    for trecho in _NEGRITO.finditer(texto):
        if trecho.start() > posicao:
            paragrafo.add_run(texto[posicao : trecho.start()])
        paragrafo.add_run(trecho.group(1)).bold = True
        posicao = trecho.end()
    if posicao < len(texto):
        paragrafo.add_run(texto[posicao:])


def _paragrafo(documento: DocumentoDocx, texto: str, estilo: str | None = None) -> None:
    paragrafo = documento.add_paragraph(style=estilo) if estilo else documento.add_paragraph()
    _adicionar_runs(paragrafo, texto)


def _fechar(documento: DocumentoDocx, pendente: list[str]) -> None:
    if pendente:
        _paragrafo(documento, " ".join(pendente))
        pendente.clear()


def markdown_para_docx(markdown: str) -> bytes:
    documento = Document()
    pendente: list[str] = []
    for linha in markdown.splitlines():
        texto = linha.strip()
        if not texto:
            _fechar(documento, pendente)
            continue
        if texto.startswith("## "):
            _fechar(documento, pendente)
            _adicionar_runs(documento.add_heading(level=2), texto[3:].strip())
        elif texto.startswith("# "):
            _fechar(documento, pendente)
            _adicionar_runs(documento.add_heading(level=1), texto[2:].strip())
        elif numerado := _NUMERADO.match(texto):
            _fechar(documento, pendente)
            _paragrafo(documento, numerado.group(2), "List Number")
        elif marcador := _MARCADOR.match(texto):
            _fechar(documento, pendente)
            _paragrafo(documento, marcador.group(1), "List Bullet")
        else:
            pendente.append(texto)
    _fechar(documento, pendente)
    buffer = BytesIO()
    documento.save(buffer)
    return buffer.getvalue()
