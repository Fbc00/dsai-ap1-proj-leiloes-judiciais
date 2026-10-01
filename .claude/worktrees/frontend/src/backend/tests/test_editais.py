import json
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

from docx import Document
from httpx import AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession

from app.agenda import service as agenda_service
from app.agenda.models import Leilao
from app.config import Settings
from app.editais.docx import markdown_para_docx
from app.editais.service import VAZIO, formatar_moeda, lance_segundo, por_extenso
from tests.test_processos import CHECKLIST_VAZIO

TIPO_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def test_formatar_moeda_padrao_brasileiro() -> None:
    assert formatar_moeda(Decimal("1234567.5")) == "R$ 1.234.567,50"
    assert formatar_moeda(Decimal("0.5")) == "R$ 0,50"
    assert formatar_moeda(Decimal("215000.00")) == "R$ 215.000,00"


def test_por_extenso_em_reais() -> None:
    texto = por_extenso(Decimal("1234.56"))
    assert "mil duzentos e trinta e quatro reais" in texto
    assert "cinquenta e seis centavos" in texto


def test_lance_segundo_e_70_por_cento_arredondando_meio_para_cima() -> None:
    assert lance_segundo(Decimal("100.00")) == Decimal("70.00")
    assert lance_segundo(Decimal("1.05")) == Decimal("0.74")
    assert lance_segundo(Decimal("215000.00")) == Decimal("150500.00")


def test_markdown_para_docx_mapeia_titulos_negrito_e_listas() -> None:
    markdown = (
        "# Título principal\n\n"
        "## Subtítulo\n\n"
        "Texto com **palavra forte** no meio\ncontinua na mesma linha lógica.\n\n"
        "1. Primeiro item\n"
        "2. Segundo item\n"
        "- Marcador\n\n"
        "**Assinatura**\n"
    )
    conteudo = markdown_para_docx(markdown)
    assert conteudo.startswith(b"PK")
    documento = Document(BytesIO(conteudo))
    paragrafos = documento.paragraphs
    estilos = [(p.style.name, p.text) for p in paragrafos]
    assert ("Heading 1", "Título principal") in estilos
    assert ("Heading 2", "Subtítulo") in estilos
    assert ("List Number", "Primeiro item") in estilos
    assert ("List Number", "Segundo item") in estilos
    assert ("List Bullet", "Marcador") in estilos
    corpo = next(p for p in paragrafos if p.text.startswith("Texto com"))
    assert corpo.text == "Texto com palavra forte no meio continua na mesma linha lógica."
    assert [r.text for r in corpo.runs if r.bold] == ["palavra forte"]
    assinatura = next(p for p in paragrafos if p.text == "Assinatura")
    assert all(r.bold for r in assinatura.runs)


async def test_gerar_edital_preenche_dados_do_checklist_e_do_leilao(
    cliente_logado: AsyncClient, leilao_agendado: Leilao, settings_teste: Settings
) -> None:
    resposta = await cliente_logado.post(
        "/api/editais", json={"leilao_id": str(leilao_agendado.id)}
    )
    assert resposta.status_code == 201, resposta.text
    corpo = resposta.json()
    assert corpo["leilao_id"] == str(leilao_agendado.id)
    assert corpo["processo_id"] == str(leilao_agendado.processo_id)
    assert corpo["atualizado_em"] is None
    md: str = corpo["conteudo_markdown"]
    primeiro = leilao_agendado.primeiro_leilao_em.astimezone(settings_teste.tz)
    segundo = leilao_agendado.segundo_leilao_em.astimezone(settings_teste.tz)
    assert f"**Primeiro Leilão**: {primeiro:%d/%m/%Y} às {primeiro:%H:%M}h." in md
    assert f"**Segundo Leilão**: {segundo:%d/%m/%Y} às {segundo:%H:%M}h." in md
    assert "CLAUDIO CEZAR CAVALCANTES" in md
    assert "2ª Vara Federal Cível e Criminal da Comarca de Araguaína - TO" in md
    assert "1003966-15.2022.4.01.4301" in md
    assert "Matrícula nº 6.351" in md
    assert "Matrícula nº 6.235" in md
    assert "R$ 250.000,00 (duzentos e cinquenta mil reais" in md
    assert "**Lance Inicial em 1º Leilão**: R$ 215.000,00" in md
    assert "**Lance Inicial em 2º Leilão**: R$ 150.500,00" in md
    assert "www.norteleiloes.com.br" in md
    assert "Sandro de Oliveira" in md
    assert "Eurivaldo Soares de Andrade (Id 2221536095)" in md


async def test_edital_com_checklist_vazio_usa_tracos(
    cliente_logado: AsyncClient, leilao_agendado: Leilao
) -> None:
    vazio = json.loads(json.dumps(CHECKLIST_VAZIO))
    atualizado = await cliente_logado.put(
        f"/api/processos/{leilao_agendado.processo_id}/checklist", json=vazio
    )
    assert atualizado.status_code == 200
    resposta = await cliente_logado.post(
        "/api/editais", json={"leilao_id": str(leilao_agendado.id)}
    )
    assert resposta.status_code == 201, resposta.text
    md: str = resposta.json()["conteudo_markdown"]
    assert md.count(VAZIO) >= 5
    assert f"**Valor da execução:** {VAZIO}." in md
    assert f"**Lance Inicial em 1º Leilão**: {VAZIO}" in md


async def test_edital_leilao_cancelado_409(
    cliente_logado: AsyncClient, leilao_agendado: Leilao, session: AsyncSession
) -> None:
    await agenda_service.cancelar(session, leilao_agendado.id)
    resposta = await cliente_logado.post(
        "/api/editais", json={"leilao_id": str(leilao_agendado.id)}
    )
    assert resposta.status_code == 409
    assert resposta.json() == {"detail": "Leilão cancelado não gera edital"}


async def test_edital_leilao_inexistente_404(cliente_logado: AsyncClient) -> None:
    resposta = await cliente_logado.post("/api/editais", json={"leilao_id": str(uuid4())})
    assert resposta.status_code == 404
    assert (await cliente_logado.get(f"/api/editais/{uuid4()}")).status_code == 404


async def test_download_markdown_com_content_disposition(
    cliente_logado: AsyncClient, leilao_agendado: Leilao
) -> None:
    criado = (
        await cliente_logado.post("/api/editais", json={"leilao_id": str(leilao_agendado.id)})
    ).json()
    resposta = await cliente_logado.get(f"/api/editais/{criado['id']}/download")
    assert resposta.status_code == 200
    assert resposta.headers["content-type"].startswith("text/markdown")
    assert (
        resposta.headers["content-disposition"]
        == 'attachment; filename="edital-1003966-15.2022.4.01.4301.md"'
    )
    assert resposta.text == criado["conteudo_markdown"]


async def test_download_docx_com_juiz_em_negrito(
    cliente_logado: AsyncClient, leilao_agendado: Leilao
) -> None:
    criado = (
        await cliente_logado.post("/api/editais", json={"leilao_id": str(leilao_agendado.id)})
    ).json()
    resposta = await cliente_logado.get(
        f"/api/editais/{criado['id']}/download", params={"formato": "docx"}
    )
    assert resposta.status_code == 200
    assert resposta.headers["content-type"].startswith(TIPO_DOCX)
    assert (
        resposta.headers["content-disposition"]
        == 'attachment; filename="edital-1003966-15.2022.4.01.4301.docx"'
    )
    assert resposta.content.startswith(b"PK")
    documento = Document(BytesIO(resposta.content))
    assert any(
        p.style.name == "Heading 1" and p.text == "HASTA PÚBLICA" for p in documento.paragraphs
    )
    runs_negrito = [r.text for p in documento.paragraphs for r in p.runs if r.bold]
    assert "CLAUDIO CEZAR CAVALCANTES" in runs_negrito
    assert any(p.style.name == "List Number" for p in documento.paragraphs)


async def test_download_formato_invalido_422(
    cliente_logado: AsyncClient, leilao_agendado: Leilao
) -> None:
    criado = (
        await cliente_logado.post("/api/editais", json={"leilao_id": str(leilao_agendado.id)})
    ).json()
    resposta = await cliente_logado.get(
        f"/api/editais/{criado['id']}/download", params={"formato": "pdf"}
    )
    assert resposta.status_code == 422


async def test_put_edital_atualiza_conteudo_e_download_reflete(
    cliente_logado: AsyncClient, leilao_agendado: Leilao
) -> None:
    criado = (
        await cliente_logado.post("/api/editais", json={"leilao_id": str(leilao_agendado.id)})
    ).json()
    novo = criado["conteudo_markdown"].replace("HASTA PÚBLICA", "HASTA PÚBLICA REVISADA")
    resposta = await cliente_logado.put(
        f"/api/editais/{criado['id']}", json={"conteudo_markdown": novo}
    )
    assert resposta.status_code == 200
    assert resposta.json()["id"] == criado["id"]
    assert resposta.json()["conteudo_markdown"] == novo
    assert resposta.json()["atualizado_em"] is not None
    download = await cliente_logado.get(f"/api/editais/{criado['id']}/download")
    assert "HASTA PÚBLICA REVISADA" in download.text
    assert (await cliente_logado.get(f"/api/editais/{criado['id']}")).json()[
        "conteudo_markdown"
    ] == novo


async def test_put_edital_vazio_422(cliente_logado: AsyncClient, leilao_agendado: Leilao) -> None:
    criado = (
        await cliente_logado.post("/api/editais", json={"leilao_id": str(leilao_agendado.id)})
    ).json()
    assert (
        await cliente_logado.put(f"/api/editais/{criado['id']}", json={"conteudo_markdown": ""})
    ).status_code == 422
    assert (
        await cliente_logado.put(
            f"/api/editais/{criado['id']}", json={"conteudo_markdown": "   \n"}
        )
    ).status_code == 422
    assert (await cliente_logado.get(f"/api/editais/{criado['id']}")).json()[
        "conteudo_markdown"
    ] == criado["conteudo_markdown"]


async def test_edital_id_do_leilao_aponta_para_o_mais_recente(
    cliente_logado: AsyncClient, leilao_agendado: Leilao
) -> None:
    primeiro = (
        await cliente_logado.post("/api/editais", json={"leilao_id": str(leilao_agendado.id)})
    ).json()
    segundo = (
        await cliente_logado.post("/api/editais", json={"leilao_id": str(leilao_agendado.id)})
    ).json()
    assert primeiro["id"] != segundo["id"]
    detalhe = (await cliente_logado.get(f"/api/processos/{leilao_agendado.processo_id}")).json()
    assert detalhe["leilao"]["edital_id"] == segundo["id"]
    get = await cliente_logado.get(f"/api/editais/{segundo['id']}")
    assert get.status_code == 200
    assert get.json()["conteudo_markdown"] == segundo["conteudo_markdown"]


async def test_download_com_numero_processo_nao_ascii_sanitiza_nome(
    cliente_logado: AsyncClient, leilao_agendado: Leilao
) -> None:
    detalhe = (await cliente_logado.get(f"/api/processos/{leilao_agendado.processo_id}")).json()
    checklist = detalhe["checklist"]
    checklist["numero_processo"] = 'Nº 1003966–15 "x"'
    assert (
        await cliente_logado.put(
            f"/api/processos/{leilao_agendado.processo_id}/checklist", json=checklist
        )
    ).status_code == 200
    criado = (
        await cliente_logado.post("/api/editais", json={"leilao_id": str(leilao_agendado.id)})
    ).json()
    for formato in ("md", "docx"):
        resposta = await cliente_logado.get(
            f"/api/editais/{criado['id']}/download", params={"formato": formato}
        )
        assert resposta.status_code == 200
        assert (
            resposta.headers["content-disposition"]
            == f'attachment; filename="edital-N_1003966_15_x_.{formato}"'
        )
