import ssl
from email import message_from_bytes
from email.policy import default
from io import BytesIO
from uuid import uuid4

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.agenda import service as agenda_service
from app.agenda.models import Leilao
from app.config import Settings
from app.email import sender
from app.email.sender import EnvioEmailFalhou
from app.marketing import service as marketing_service
from app.marketing.service import COLUNAS

COLUNAS_ESPERADAS: list[str] = [
    "Nº Processo",
    "Vara",
    "Executado(s)",
    "Bem",
    "Matrícula",
    "Avaliação (R$)",
    "Lance 1º (R$)",
    "Lance 2º (R$)",
    "1º Leilão",
    "2º Leilão",
    "Edital",
]


def _planilha(conteudo: bytes) -> tuple[list[str], list[list[object]]]:
    aba = load_workbook(BytesIO(conteudo))["Leilões"]
    linhas = [[c.value for c in linha] for linha in aba.iter_rows()]
    return [str(v) for v in linhas[0]], linhas[1:]


def test_colunas_da_planilha() -> None:
    assert COLUNAS == COLUNAS_ESPERADAS


async def test_planilha_uma_linha_por_bem_de_leilao_agendado(
    cliente_logado: AsyncClient, leilao_agendado: Leilao
) -> None:
    resposta = await cliente_logado.get("/api/marketing/planilha")
    assert resposta.status_code == 200
    assert resposta.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert resposta.headers["content-disposition"] == 'attachment; filename="leiloes.xlsx"'
    header, linhas = _planilha(resposta.content)
    assert header == COLUNAS_ESPERADAS
    assert len(linhas) == 2
    assert linhas[0][0] == "1003966-15.2022.4.01.4301"
    assert linhas[0][4] == "6.351"
    assert linhas[1][4] == "6.235"
    assert linhas[0][5] == 120000.0
    assert linhas[0][6] == 120000.0
    assert linhas[0][7] == 84000.0
    assert linhas[0][10] == "Não"


async def test_planilha_ignora_leilao_cancelado(
    cliente_logado: AsyncClient, leilao_agendado: Leilao, session: AsyncSession
) -> None:
    await agenda_service.cancelar(session, leilao_agendado.id)
    _, linhas = _planilha((await cliente_logado.get("/api/marketing/planilha")).content)
    assert linhas == []


async def test_envio_sem_edital_409(cliente_logado: AsyncClient, leilao_agendado: Leilao) -> None:
    resposta = await cliente_logado.post(
        "/api/marketing/envios", json={"leilao_id": str(leilao_agendado.id)}
    )
    assert resposta.status_code == 409
    assert resposta.json() == {"detail": "Gere o edital antes de enviar"}


async def test_envio_leilao_inexistente_404(cliente_logado: AsyncClient) -> None:
    resposta = await cliente_logado.post("/api/marketing/envios", json={"leilao_id": str(uuid4())})
    assert resposta.status_code == 404


async def test_envio_ok_grava_eml_com_anexo(
    cliente_logado: AsyncClient, leilao_agendado: Leilao, settings_teste: Settings
) -> None:
    await cliente_logado.post("/api/editais", json={"leilao_id": str(leilao_agendado.id)})
    resposta = await cliente_logado.post(
        "/api/marketing/envios", json={"leilao_id": str(leilao_agendado.id)}
    )
    assert resposta.status_code == 201, resposta.text
    corpo = resposta.json()
    assert corpo["status"] == "enviado"
    assert corpo["erro"] is None
    assert corpo["destinatarios"] == ["marketing@exemplo.com.br"]
    assert corpo["numero_processo"] == "1003966-15.2022.4.01.4301"

    arquivos = sorted((settings_teste.MEDIA_ROOT / "outbox").glob("*.eml"))
    assert len(arquivos) == 1
    mensagem = message_from_bytes(arquivos[0].read_bytes(), policy=default)
    assert mensagem["To"] == "marketing@exemplo.com.br"
    assert mensagem["From"] == "leiloes@exemplo.com.br"
    assert mensagem["Subject"] == "Leilão agendado – processo 1003966-15.2022.4.01.4301"
    anexos = [parte.get_filename() for parte in mensagem.iter_attachments()]
    assert anexos == ["leiloes.xlsx"]
    header, linhas = _planilha(next(mensagem.iter_attachments()).get_payload(decode=True))
    assert header == COLUNAS_ESPERADAS
    assert linhas[0][10] == "Sim"


async def test_envio_falha_smtp_grava_erro_e_502(
    cliente_logado: AsyncClient, leilao_agendado: Leilao, monkeypatch: pytest.MonkeyPatch
) -> None:
    await cliente_logado.post("/api/editais", json={"leilao_id": str(leilao_agendado.id)})

    async def _falha(*args: object, **kwargs: object) -> None:
        raise EnvioEmailFalhou("conexão recusada")

    monkeypatch.setattr(marketing_service, "send_email", _falha)
    resposta = await cliente_logado.post(
        "/api/marketing/envios", json={"leilao_id": str(leilao_agendado.id)}
    )
    assert resposta.status_code == 502
    assert resposta.json() == {"detail": "Falha ao enviar e-mail"}
    envios = (await cliente_logado.get("/api/marketing/envios")).json()
    assert len(envios) == 1
    assert envios[0]["status"] == "erro"
    assert envios[0]["erro"] == "conexão recusada"


async def test_listar_envios_mais_recente_primeiro(
    cliente_logado: AsyncClient, leilao_agendado: Leilao
) -> None:
    await cliente_logado.post("/api/editais", json={"leilao_id": str(leilao_agendado.id)})
    primeiro = (
        await cliente_logado.post(
            "/api/marketing/envios", json={"leilao_id": str(leilao_agendado.id)}
        )
    ).json()
    segundo = (
        await cliente_logado.post(
            "/api/marketing/envios", json={"leilao_id": str(leilao_agendado.id)}
        )
    ).json()
    envios = (await cliente_logado.get("/api/marketing/envios")).json()
    assert [e["id"] for e in envios] == [segundo["id"], primeiro["id"]]


async def test_envio_leilao_cancelado_409(
    cliente_logado: AsyncClient, leilao_agendado: Leilao, session: AsyncSession
) -> None:
    await cliente_logado.post("/api/editais", json={"leilao_id": str(leilao_agendado.id)})
    await agenda_service.cancelar(session, leilao_agendado.id)
    resposta = await cliente_logado.post(
        "/api/marketing/envios", json={"leilao_id": str(leilao_agendado.id)}
    )
    assert resposta.status_code == 409
    assert resposta.json() == {"detail": "Leilão cancelado não é enviado ao marketing"}
    assert (await cliente_logado.get("/api/marketing/envios")).json() == []


async def test_smtp_starttls_verifica_certificado(
    settings_teste: Settings, monkeypatch: pytest.MonkeyPatch
) -> None:
    contextos: list[object] = []

    class SmtpFalso:
        def __init__(self, *args: object, **kwargs: object) -> None:
            pass

        def __enter__(self) -> "SmtpFalso":
            return self

        def __exit__(self, *args: object) -> None:
            pass

        def starttls(self, context: object = None) -> None:
            contextos.append(context)

        def login(self, usuario: str, senha: str) -> None:
            pass

        def send_message(self, mensagem: object) -> None:
            pass

    monkeypatch.setattr(sender.smtplib, "SMTP", SmtpFalso)
    settings = settings_teste.model_copy(
        update={"EMAIL_BACKEND": "smtp", "SMTP_HOST": "smtp.exemplo", "SMTP_TLS": True}
    )
    await sender.send_email(settings, to=["a@b.c"], subject="s", body="b", attachment=None)
    assert len(contextos) == 1
    contexto = contextos[0]
    assert isinstance(contexto, ssl.SSLContext)
    assert contexto.verify_mode == ssl.CERT_REQUIRED
    assert contexto.check_hostname is True
